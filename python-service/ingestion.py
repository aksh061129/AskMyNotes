"""Document ingestion pipeline: parse → chunk → embed → FAISS index."""

import json
import fitz  # PyMuPDF
import faiss
import numpy as np
from pathlib import Path
from docx import Document
from pptx import Presentation
from PIL import Image
import pytesseract
import zipfile
import tempfile
import os

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

from config import DATA_DIR, CHUNK_SIZE, CHUNK_OVERLAP
from embedder import embedder
from llama_index.core.node_parser import SentenceSplitter


# ── Metadata sidecar per subject ──────────────────────────────────────

def _subject_dir(subject_id: str) -> Path:
    d = DATA_DIR / subject_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def _load_metadata(subject_id: str) -> list[dict]:
    path = _subject_dir(subject_id) / "chunks_meta.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return []


def _save_metadata(subject_id: str, meta: list[dict]):
    path = _subject_dir(subject_id) / "chunks_meta.json"
    path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def _load_faiss_index(subject_id: str) -> faiss.IndexFlatIP | None:
    path = _subject_dir(subject_id) / "index.faiss"
    if path.exists():
        return faiss.read_index(str(path))
    return None


def _save_faiss_index(subject_id: str, index: faiss.IndexFlatIP):
    path = _subject_dir(subject_id) / "index.faiss"
    faiss.write_index(index, str(path))


# ── Document Parsing ──────────────────────────────────────────────────

def _parse_pdf(file_path: str) -> list[dict]:
    """Parse PDF → list of {text, page, filename}."""
    pages = []
    doc = fitz.open(file_path)
    filename = Path(file_path).name
    for page_num, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if text:
            pages.append({
                "text": text,
                "page": page_num,
                "filename": filename,
            })
    doc.close()
    return pages


def _parse_txt(file_path: str) -> list[dict]:
    """Parse TXT → single entry."""
    text = Path(file_path).read_text(encoding="utf-8", errors="ignore").strip()
    if not text:
        return []
    return [{
        "text": text,
        "page": 1,
        "filename": Path(file_path).name,
    }]


def _parse_docx(file_path: str) -> list[dict]:
    """Parse DOCX/Word document → list of text entries."""
    doc = Document(file_path)
    pages = []
    filename = Path(file_path).name

    # Extract normal paragraphs
    for paragraph_index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            pages.append({
                "text": text,
                "page": None,
                "filename": filename,
            })

    # Extract tables
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                rows.append(row_text)

        if rows:
            pages.append({
                "text": "\n".join(rows),
                "page": None,
                "filename": filename,
            })

    return pages


def _parse_pptx(file_path: str) -> list[dict]:
    """Parse PPTX/PowerPoint presentation into text entries."""
    presentation = Presentation(file_path)
    slides = []
    filename = Path(file_path).name

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    texts.append(text)

        if texts:
            slides.append({
                "text": "\n".join(texts),
                "page": slide_number,
                "filename": filename,
            })

    return slides


def _parse_image(file_path: str) -> list[dict]:
    """Parse image (PNG, JPG, JPEG, WEBP) using pytesseract OCR → single entry."""
    image = Image.open(file_path)
    text = pytesseract.image_to_string(image).strip()
    if not text:
        return []
    return [{
        "text": text,
        "page": 1,
        "filename": Path(file_path).name,
    }]

def _parse_supported_file(file_path: str) -> list[dict]:
    """Parse one supported file using the existing parsers."""
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext in (".txt", ".text"):
        return _parse_txt(file_path)
    elif ext == ".docx":
        return _parse_docx(file_path)
    elif ext == ".pptx":
        return _parse_pptx(file_path)
    elif ext in (".png", ".jpg", ".jpeg", ".webp"):
        return _parse_image(file_path)
    else:
        return []
def _parse_zip(file_path: str) -> list[dict]:
    """Extract and parse supported files from a ZIP archive."""
    supported_extensions = {
        ".pdf",
        ".txt",
        ".text",
        ".docx",
        ".pptx",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    }

    entries = []

    with tempfile.TemporaryDirectory() as extract_dir:
        with zipfile.ZipFile(file_path, "r") as zip_ref:
            zip_ref.extractall(extract_dir)

        for root, _, files in os.walk(extract_dir):
            for filename in files:
                ext = Path(filename).suffix.lower()

                if ext not in supported_extensions:
                    continue

                extracted_path = os.path.join(root, filename)

                try:
                    parsed = _parse_supported_file(extracted_path)
                    entries.extend(parsed)
                except Exception as e:
                    print(f"Skipping {filename}: {e}")

    return entries

# ── Chunking ──────────────────────────────────────────────────────────

def _chunk_documents(pages: list[dict]) -> list[dict]:
    """Chunk page texts using LlamaIndex SentenceSplitter, preserving metadata."""
    splitter = SentenceSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    chunks = []

    for page_info in pages:
        text = page_info["text"]
        split_texts = splitter.split_text(text)
        for chunk_text in split_texts:
            chunks.append({
                "text": chunk_text,
                "filename": page_info["filename"],
                "page": page_info["page"],
            })

    return chunks


# ── Public API ────────────────────────────────────────────────────────

def ingest_document(subject_id: str, file_path: str) -> dict:
    """
    Full ingestion pipeline for a single document.

    Returns summary dict with chunk count and status.
    """
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        pages = _parse_pdf(file_path)
    elif ext in (".txt", ".text"):
        pages = _parse_txt(file_path)
    elif ext == ".docx":
        pages = _parse_docx(file_path)
    elif ext == ".pptx":
        pages = _parse_pptx(file_path)
    elif ext == ".zip":
        pages = _parse_pptx(file_path)
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".zip"):
        pages = _parse_image(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not pages:
        return {"status": "empty", "chunks_added": 0}

    # Chunk
    chunks = _chunk_documents(pages)
    if not chunks:
        return {"status": "empty", "chunks_added": 0}

    # Embed
    texts = [c["text"] for c in chunks]
    vectors = embedder.encode(texts)

    # Normalise for cosine similarity via inner product
    faiss.normalize_L2(vectors)

    # Load or create FAISS index
    index = _load_faiss_index(subject_id)
    if index is None:
        index = faiss.IndexFlatIP(embedder.dimension)  # Inner product on normed vecs = cosine sim
    index.add(vectors)

    # Update metadata
    existing_meta = _load_metadata(subject_id)
    base_idx = len(existing_meta)
    for i, chunk in enumerate(chunks):
        chunk["chunk_index"] = base_idx + i
    existing_meta.extend(chunks)

    # Persist
    _save_faiss_index(subject_id, index)
    _save_metadata(subject_id, existing_meta)

    return {
        "status": "success",
        "chunks_added": len(chunks),
        "total_chunks": len(existing_meta),
        "filename": Path(file_path).name,
    }


def get_subject_stats(subject_id: str) -> dict:
    """Return stats about a subject's index."""
    meta = _load_metadata(subject_id)
    index = _load_faiss_index(subject_id)
    files = list({m["filename"] for m in meta})
    return {
        "subject_id": subject_id,
        "total_chunks": len(meta),
        "total_vectors": index.ntotal if index else 0,
        "files": files,
    }
